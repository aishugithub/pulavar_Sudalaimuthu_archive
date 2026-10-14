"""Build highlighted .docx and .html review copies from a transcription .txt.

Markers in the source text:
  {?word?}   uncertain reading -> yellow highlight
  {note}     transcriber's note -> grey italic
  [பக்கம் N] page marker from the scan
  # / ##     title / section heading
"""
import html, re, sys
from pathlib import Path
from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

TOKEN = re.compile(r'(\{\?.*?\?\}|\{[^?].*?\})')
FONT = 'Noto Sans Tamil'


def blocks(text):
    """Split into paragraphs (blank-line separated); keep line breaks inside verse."""
    for chunk in re.split(r'\n\s*\n', text.strip()):
        yield chunk.strip('\n')


def flags(text):
    page, out = None, []
    for line in text.splitlines():
        m = re.match(r'\[பக்கம் (\d+)\]', line)
        if m:
            page = m.group(1)
        for w in re.findall(r'\{\?(.*?)\?\}', line):
            out.append((page, w))
    return out


def set_font(run, size=None):
    run.font.name = FONT
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = rpr.makeelement(qn('w:rFonts'), {})
        rpr.append(rfonts)
    for attr in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rfonts.set(qn(attr), FONT)
    if size:
        run.font.size = Pt(size)


def add_runs(par, line, size=12):
    for part in TOKEN.split(line):
        if not part:
            continue
        if part.startswith('{?'):
            r = par.add_run(part[2:-2] + ' (?)')
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        elif part.startswith('{'):
            r = par.add_run('[' + part[1:-1] + ']')
            r.italic = True
            r.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
        else:
            r = par.add_run(part)
        set_font(r, size)


def build_docx(text, title, out):
    doc = Document()
    st = doc.styles['Normal']
    st.font.name = FONT
    st.font.size = Pt(12)
    note = doc.add_paragraph()
    add_runs(note, 'மஞ்சள் நிறத்தில் உள்ள சொற்கள் (?) — கையெழுத்தில் தெளிவாகப் படிக்க இயலாதவை; ஸ்கேனுடன் ஒப்பிட்டுச் சரிபார்க்கவும்.', 10)
    for b in blocks(text):
        if b.startswith('# '):
            h = doc.add_heading(level=0); add_runs(h, b[2:], 20)
        elif b.startswith('## '):
            h = doc.add_heading(level=2); add_runs(h, b[3:], 14)
        elif re.fullmatch(r'\[.*\]', b):
            p = doc.add_paragraph(); r = p.add_run(b); set_font(r, 9)
            r.font.color.rgb = RGBColor(0x99, 0x66, 0x00)
        else:
            p = doc.add_paragraph()
            for i, line in enumerate(b.split('\n')):
                if i:
                    p.add_run().add_break()
                add_runs(p, line)
    fl = flags(text)
    doc.add_page_break()
    h = doc.add_heading(level=2); add_runs(h, f'சரிபார்க்க வேண்டிய சொற்கள் ({len(fl)})', 14)
    for pg, w in fl:
        p = doc.add_paragraph(style='List Number'); add_runs(p, f'பக்கம் {pg}: {{?{w}?}}', 11)
    doc.save(out)


def inline_html(line):
    out = []
    for part in TOKEN.split(line):
        if not part:
            continue
        if part.startswith('{?'):
            out.append(f'<mark>{html.escape(part[2:-2])}<sup>?</sup></mark>')
        elif part.startswith('{'):
            out.append(f'<span class="note">[{html.escape(part[1:-1])}]</span>')
        else:
            out.append(html.escape(part))
    return ''.join(out)


def build_html(text, title, out):
    body = []
    for b in blocks(text):
        if b.startswith('# '):
            body.append(f'<h1>{inline_html(b[2:])}</h1>')
        elif b.startswith('## '):
            body.append(f'<h2>{inline_html(b[3:])}</h2>')
        elif re.fullmatch(r'\[.*\]', b):
            body.append(f'<p class="page">{html.escape(b)}</p>')
        else:
            body.append('<p>' + '<br>'.join(inline_html(l) for l in b.split('\n')) + '</p>')
    fl = flags(text)
    items = ''.join(f'<li>பக்கம் {pg}: <mark>{html.escape(w)}</mark></li>' for pg, w in fl)
    page = f'''<!doctype html><html lang="ta"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>{html.escape(title)}</title>
<link href="https://fonts.googleapis.com/css2?family=Noto+Sans+Tamil:wght@400;700&display=swap" rel="stylesheet">
<style>
:root{{--bg:#fbf8f1;--ink:#2b2420;--soft:#7a6f66;--mark:#ffe766;--acc:#7a2e2e}}
@media (prefers-color-scheme:dark){{:root{{--bg:#1d1a17;--ink:#ece6dc;--soft:#a89c90;--mark:#7a6400;--acc:#e0a0a0}}}}
body{{background:var(--bg);color:var(--ink);font-family:'Noto Sans Tamil',sans-serif;line-height:1.9;
max-width:720px;margin:0 auto;padding:16px}}
h1{{color:var(--acc);text-align:center}} h2{{color:var(--acc);margin-top:2em}}
mark{{background:var(--mark);color:inherit;padding:0 2px;border-radius:3px}}
.page{{color:var(--soft);font-size:.8em;border-top:1px dashed var(--soft);padding-top:.4em;margin-top:2em}}
.note{{color:var(--soft);font-style:italic}}
.help{{background:rgba(127,127,127,.12);padding:10px 12px;border-radius:8px;font-size:.9em}}
</style></head><body>
<p class="help">மஞ்சள் நிறச் சொற்கள் <mark>இப்படி<sup>?</sup></mark> — கையெழுத்தில் உறுதியாகப் படிக்க இயலாதவை. கீழே முழுப் பட்டியல் உள்ளது.</p>
{''.join(body)}
<h2>சரிபார்க்க வேண்டிய சொற்கள் ({len(fl)})</h2><ol>{items}</ol>
</body></html>'''
    Path(out).write_text(page, encoding='utf-8')


if __name__ == '__main__':
    src = Path(sys.argv[1]); title = sys.argv[2]
    text = src.read_text(encoding='utf-8')
    build_docx(text, title, src.with_suffix('.docx'))
    build_html(text, title, src.with_suffix('.html'))
    print(len(flags(text)), 'flagged words')
